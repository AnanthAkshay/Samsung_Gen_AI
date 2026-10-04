import shutil
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

def update_slides():
    src_file = "MSRIT_Cache_Me.pptx"
    backup_file = "MSRIT_Cache_Me_backup.pptx"
    shutil.copyfile(src_file, backup_file)
    print(f"Created backup at {backup_file}")

    prs = Presentation(src_file)
    slide8 = prs.slides[7]

    # Shape 3: Strict Pass Rate
    s3 = slide8.shapes[3]
    s3.text_frame.paragraphs[0].runs[0].text = "45.0%"
    s3.text_frame.paragraphs[1].runs[0].text = "Strict Pass Rate (Exact Match)"

    # Shape 4: Tool Selection F1
    s4 = slide8.shapes[4]
    s4.text_frame.paragraphs[0].runs[0].text = "0.823"
    s4.text_frame.paragraphs[1].runs[0].text = "Tool Selection F1 (vs 0.786 baseline)"

    # Shape 5: Precision
    s5 = slide8.shapes[5]
    s5.text_frame.paragraphs[0].runs[0].text = "0.835"
    s5.text_frame.paragraphs[1].runs[0].text = "Precision"

    # Shape 6: Recall
    s6 = slide8.shapes[6]
    s6.text_frame.paragraphs[0].runs[0].text = "0.838"
    s6.text_frame.paragraphs[1].runs[0].text = "Recall"

    # Shape 7: Average Response Latency
    s7 = slide8.shapes[7]
    s7.text_frame.paragraphs[0].runs[0].text = "11.73 s"
    s7.text_frame.paragraphs[1].runs[0].text = "Average Response Latency"

    # Shape 8: 100 / 100 Verified
    s8 = slide8.shapes[8]
    s8.text_frame.paragraphs[0].runs[0].text = "100 / 100"
    s8.text_frame.paragraphs[1].runs[0].text = "Full Benchmark Verified"

    # Shape 9: Baseline citation
    s9 = slide8.shapes[9]
    s9.text_frame.paragraphs[0].runs[0].text = "Published Baselines (arXiv:2604.04847): GPT-Realtime (F1: 0.876) | Gemini 2.5 (F1: 0.786) | Cascaded (10.12s)"

    # Shape 10: Limitations & Comparison
    s10 = slide8.shapes[10]
    tf10 = s10.text_frame
    tf10.paragraphs[0].runs[0].text = "BENCHMARK COMPARISON & ARCHITECTURAL UPGRADES"
    tf10.paragraphs[1].runs[0].text = "100% Turn-Taking: Zero silent drops across all 100 scenarios (vs 8%-22% dropped in published Gemini baselines)"
    tf10.paragraphs[2].runs[0].text = "Exceeds Gemini 2.5 Baseline F1: 0.823 vs 0.786; lowest interruption rate at 5.0% (vs 14.1% Gemini, 13.5% GPT)"
    tf10.paragraphs[3].runs[0].text = "Explicit State & Idempotency: SessionStateManager tracks slots, cancels superseded calls, and blocks duplicate mutations"
    tf10.paragraphs[4].runs[0].text = "Non-Blocking Async Engine: run_in_executor and startup schema preloading eliminate event-loop audio stalls"

    prs.save(src_file)
    print("MSRIT_Cache_Me.pptx Slide 8 successfully updated!")

if __name__ == "__main__":
    update_slides()
